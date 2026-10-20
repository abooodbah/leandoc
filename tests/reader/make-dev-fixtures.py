"""Generate the reader's development fixtures.

Usage: python3 tests/reader/make-dev-fixtures.py [output-dir]

Writes small deterministic .docx files into tests/reader/dev-fixtures/ (fixed
ZIP timestamps and core properties) plus reports.json, a static report per
file in the shape of the core's ReportToJson(). The outline and paragraph
indexes follow SPEC section 3.6 so the reader tests can run before the core
is merged. Requires python-docx 1.2 and Pillow.
"""

from __future__ import annotations

import datetime
import io
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FIXED_DATE = datetime.datetime(2026, 1, 2, 3, 4, 5)
ZIP_DATE = (2026, 1, 2, 3, 4, 6)


def new_document(title: str) -> Document:
    document = Document()
    props = document.core_properties
    props.title = title
    props.author = "LeanDoc tests"
    props.last_modified_by = "LeanDoc tests"
    props.created = FIXED_DATE
    props.modified = FIXED_DATE
    props.last_printed = FIXED_DATE
    props.revision = 1
    props.comments = ""
    props.keywords = ""
    props.subject = ""
    props.category = ""
    return document


def page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_bidi(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(parse_xml(f"<w:bidi {nsdecls('w')}/>"))


def rtl_run(paragraph, text: str):
    run = paragraph.add_run(text)
    rpr = run._r.get_or_add_rPr()
    rpr.append(parse_xml(f"<w:rtl {nsdecls('w')}/>"))
    return run


def set_numbering(paragraph, num_id: int, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/>'
        f'<w:numId w:val="{num_id}"/></w:numPr>'))


def add_abstract_numbering(document: Document, abstract_id: int, num_id: int,
                           levels: list[tuple[str, str, str | None]]) -> None:
    """levels: (numFmt, lvlText, font or None)."""
    numbering = document.part.numbering_part.element
    parts = [f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">',
             '<w:multiLevelType w:val="hybridMultilevel"/>']
    for index, (fmt, text, font) in enumerate(levels):
        left = 720 * (index + 1)
        rpr = ""
        if font:
            rpr = (f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" '
                   f'w:hint="default"/></w:rPr>')
        text_xml = text.replace("&", "&amp;").replace('"', "&quot;")
        parts.append(
            f'<w:lvl w:ilvl="{index}"><w:start w:val="1"/>'
            f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text_xml}"/>'
            f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{left}" '
            f'w:hanging="360"/></w:pPr>{rpr}</w:lvl>')
    parts.append("</w:abstractNum>")
    abstract = parse_xml("".join(parts))
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}">'
        f'<w:abstractNumId w:val="{abstract_id}"/></w:num>'))


def add_hyperlink(paragraph, url: str | None, text: str, anchor: str | None = None):
    attributes = ""
    if url is not None:
        rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        attributes += f' r:id="{rel_id}"'
    if anchor is not None:
        attributes += f' w:anchor="{anchor}"'
    paragraph._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")}{attributes}><w:r><w:rPr>'
        f'<w:color w:val="1D4ED8"/><w:u w:val="single"/></w:rPr>'
        f'<w:t xml:space="preserve">{text}</w:t></w:r></w:hyperlink>'))


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, parse_xml(
        f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{name}"/>'))
    paragraph._p.append(parse_xml(
        f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>'))


def add_sym(paragraph, font: str, char: str) -> None:
    paragraph._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:sym w:font="{font}" w:char="{char}"/></w:r>'))


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


# --- documents -------------------------------------------------------------


def make_headings(document: Document) -> None:
    document.add_paragraph("Quarterly field report", style="Title")
    document.add_paragraph(
        "This report collects the observations from the spring survey. "
        "It is a fixture for the LeanDoc reader tests.")
    document.add_heading("Introduction", level=1)
    document.add_paragraph(
        "The survey covered twelve sites along the river. Each site was "
        "visited twice, once in April and once in May.")
    document.add_heading("Method", level=1)
    document.add_paragraph("Teams used the same checklist at every site.")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Site"
    table.cell(0, 1).text = "Visits"
    table.cell(1, 0).text = "North bank"
    table.cell(1, 1).text = "2"
    document.add_heading("Sampling", level=2)
    document.add_paragraph(
        "Samples were taken at dawn. The marker word appears here: "
        "kingfisher.")
    page_break(document)
    document.add_heading("Results", level=1)
    for index in range(8):
        document.add_paragraph(
            f"Result paragraph {index + 1}. The water was clear and the "
            "kingfisher was seen near the old mill.")
    document.add_heading("Water quality", level=2)
    document.add_paragraph("Nitrate levels stayed within the expected range.")
    document.add_heading("Detail", level=3)
    document.add_paragraph("Measurements were repeated for each sample.")
    page_break(document)
    document.add_heading("Conclusion", level=1)
    document.add_paragraph("The closing paragraph mentions the kingfisher once more.")


def make_bullets(document: Document) -> None:
    add_abstract_numbering(document, 90, 90, [
        ("bullet", "", "Symbol"),
        ("bullet", "o", "Courier New"),
        ("bullet", "", "Wingdings"),
    ])
    add_abstract_numbering(document, 91, 91, [
        ("bullet", "", "Wingdings"),
        ("bullet", "", "Wingdings"),
        ("bullet", "", "Wingdings"),
    ])
    add_abstract_numbering(document, 92, 92, [
        ("decimal", "%1.", None),
        ("decimal", "%1.%2.", None),
        ("lowerLetter", "%3)", None),
    ])
    document.add_heading("Bullets in symbol fonts", level=1)
    for text, level in [("First bullet in Symbol", 0),
                        ("Second level in Courier New", 1),
                        ("Third level in Wingdings", 2),
                        ("Back to the first level", 0)]:
        set_numbering(document.add_paragraph(text), 90, level)
    document.add_heading("More Wingdings bullets", level=1)
    for text, level in [("Arrow bullet item", 0),
                        ("Check mark item", 1),
                        ("Diamond item", 2)]:
        set_numbering(document.add_paragraph(text), 91, level)
    document.add_heading("Multi-level numbering", level=1)
    for text, level in [("Plan the survey", 0),
                        ("Choose the sites", 1),
                        ("Upstream site", 2),
                        ("Downstream site", 2),
                        ("Book the boats", 1),
                        ("Run the survey", 0)]:
        set_numbering(document.add_paragraph(text), 92, level)
    paragraph = document.add_paragraph("Inline symbol check: ")
    add_sym(paragraph, "Wingdings", "F0FC")
    paragraph.add_run(" done.")


def make_tables(document: Document) -> None:
    section = document.sections[0]
    section.header.paragraphs[0].text = "LeanDoc header text"
    section.footer.paragraphs[0].text = "LeanDoc footer text"
    document.add_heading("Merged table", level=1)
    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged across two columns"
    table.cell(0, 2).text = "Single"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Merged down two rows"
    table.cell(1, 1).text = "B2"
    table.cell(1, 2).text = "C2"
    table.cell(2, 1).text = "B3"
    table.cell(2, 2).text = "C3"
    table.cell(3, 0).text = "A4"
    table.cell(3, 1).text = "B4"
    table.cell(3, 2).text = "C4"
    shade(table.cell(0, 2), "F4F2EE")
    shade(table.cell(3, 0), "DBEAFE")
    document.add_paragraph("The table above has merged cells, borders and shading.")
    page_break(document)
    document.add_paragraph("Second page body text. The header repeats above.")


def make_hyperlinks(document: Document) -> None:
    document.add_heading("Links", level=1)
    paragraph = document.add_paragraph("Web link: ")
    add_hyperlink(paragraph, "https://example.com/leandoc", "Example site")
    paragraph = document.add_paragraph("Mail link: ")
    add_hyperlink(paragraph, "mailto:reader@example.com", "Write to the reader")
    paragraph = document.add_paragraph("Script link: ")
    add_hyperlink(paragraph, "javascript:alert(1)", "Script text")
    paragraph = document.add_paragraph("File link: ")
    add_hyperlink(paragraph, "file:///etc/passwd", "Local file text")
    paragraph = document.add_paragraph("Internal link: ")
    add_hyperlink(paragraph, None, "Jump to the target", anchor="leandoc_target")
    for index in range(40):
        document.add_paragraph(f"Filler paragraph {index + 1} keeps the target below the fold.")
    target = document.add_paragraph("Bookmark target paragraph.")
    add_bookmark(target, "leandoc_target", 7)
    for index in range(40):
        document.add_paragraph(f"Tail paragraph {index + 1}.")


def make_arabic(document: Document) -> None:
    add_abstract_numbering(document, 93, 93, [
        ("decimal", "%1.", None),
        ("decimal", "%1.%2.", None),
    ])
    styles = document.styles
    heading = document.add_heading("", level=1)
    set_bidi(heading)
    rtl_run(heading, "تقرير المسح")
    for text in ["هذا تقرير قصير عن المسح الميداني في فصل الربيع.",
                 "زار الفريق اثني عشر موقعا على طول النهر."]:
        paragraph = document.add_paragraph()
        set_bidi(paragraph)
        rtl_run(paragraph, text)
    # Word stores jc relative to the paragraph direction: in a right-to-left
    # paragraph "left" is the start (right) edge and "right" the end.
    for jc, text in [("left", "فقرة عند بداية السطر"), ("right", "فقرة عند نهاية السطر")]:
        paragraph = document.add_paragraph()
        set_bidi(paragraph)
        paragraph._p.get_or_add_pPr().append(parse_xml(f'<w:jc {nsdecls("w")} w:val="{jc}"/>'))
        rtl_run(paragraph, text)
    mixed = document.add_paragraph()
    set_bidi(mixed)
    rtl_run(mixed, "يعمل برنامج ")
    mixed.add_run("LeanDoc")
    rtl_run(mixed, " على ويندوز ولينكس.")
    for text, level in [("الخطوة الأولى", 0), ("تفصيل الخطوة", 1),
                        ("الخطوة الثانية", 0)]:
        paragraph = document.add_paragraph()
        set_bidi(paragraph)
        set_numbering(paragraph, 93, level)
        rtl_run(paragraph, text)
    table = document.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    tblpr = table._tbl.tblPr
    tblpr.append(parse_xml(f"<w:bidiVisual {nsdecls('w')}/>"))
    for column, text in enumerate(["الموقع", "الزيارات", "ملاحظات"]):
        cell_paragraph = table.cell(0, column).paragraphs[0]
        set_bidi(cell_paragraph)
        rtl_run(cell_paragraph, text)
    for column, text in enumerate(["الضفة الشمالية", "2", "صاف"]):
        cell_paragraph = table.cell(1, column).paragraphs[0]
        set_bidi(cell_paragraph)
        rtl_run(cell_paragraph, text)
    english = document.add_paragraph("An English paragraph after the Arabic table.")
    del styles


def make_images(document: Document, image_bytes: bytes) -> None:
    document.add_heading("Inline image", level=1)
    document.add_paragraph("The picture below is stored inside the package.")
    document.add_picture(io.BytesIO(image_bytes), width=Inches(3))
    caption = document.add_paragraph("Figure 1. A generated test picture.")
    caption.runs[0].font.size = Pt(9)
    caption.runs[0].font.color.rgb = RGBColor(0x5A, 0x55, 0x4C)


def generated_png() -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (320, 180), (244, 242, 238))
    draw = ImageDraw.Draw(image)
    for x in range(320):
        shade_value = 200 - x // 4
        draw.line([(x, 120), (x, 180)], fill=(29, 78, shade_value + 40))
    draw.ellipse([30, 30, 110, 110], fill=(29, 78, 216))
    draw.rectangle([150, 40, 290, 100], outline=(26, 25, 23), width=3)
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=False)
    return output.getvalue()


# --- deterministic output and static reports --------------------------------


def deterministic_zip(raw: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(raw))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            info = zipfile.ZipInfo(name, date_time=ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            info.create_system = 0
            target.writestr(info, source.read(name))
    return output.getvalue()


def without_parts(data: bytes, drop) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if not drop(info.filename):
                target.writestr(info, source.read(info.filename))
    return output.getvalue()


def local(element) -> str:
    return etree.QName(element).localname


def paragraph_text(paragraph) -> str:
    text = "".join(node.text or "" for node in paragraph.iter(f"{{{W}}}t")
                   if not any(local(a) == "txbxContent" for a in node.iterancestors()))
    return " ".join(text.split())


def static_report(data: bytes) -> dict:
    archive = zipfile.ZipFile(io.BytesIO(data))
    document = etree.fromstring(archive.read("word/document.xml"))
    styles = etree.fromstring(archive.read("word/styles.xml"))
    style_info = {}
    default_style = None
    for style in styles.iter(f"{{{W}}}style"):
        style_id = style.get(f"{{{W}}}styleId")
        based = style.find(f"{{{W}}}basedOn")
        level = style.find(f"{{{W}}}pPr/{{{W}}}outlineLvl")
        bidi = style.find(f"{{{W}}}pPr/{{{W}}}bidi")
        style_info[style_id] = (
            based.get(f"{{{W}}}val") if based is not None else None,
            int(level.get(f"{{{W}}}val")) if level is not None else None,
            bidi is not None)
        if style.get(f"{{{W}}}type") == "paragraph" and style.get(f"{{{W}}}default") == "1":
            default_style = style_id

    def style_level(style_id):
        seen = set()
        while style_id and style_id not in seen:
            seen.add(style_id)
            based, level, _ = style_info.get(style_id, (None, None, False))
            if level is not None:
                return level
            style_id = based
        return None

    body = document.find(f"{{{W}}}body")
    outline = []
    rtl = bool(default_style and style_info.get(default_style, (None, None, False))[2])
    first_text_seen = False
    index = 0
    for paragraph in body.iter(f"{{{W}}}p"):
        if any(local(a) == "txbxContent" for a in paragraph.iterancestors()):
            continue
        ppr = paragraph.find(f"{{{W}}}pPr")
        level = None
        if ppr is not None:
            direct = ppr.find(f"{{{W}}}outlineLvl")
            if direct is not None:
                level = int(direct.get(f"{{{W}}}val"))
            else:
                style = ppr.find(f"{{{W}}}pStyle")
                if style is not None:
                    level = style_level(style.get(f"{{{W}}}val"))
        text = paragraph_text(paragraph)
        if not first_text_seen and text:
            first_text_seen = True
            if ppr is not None and ppr.find(f"{{{W}}}bidi") is not None:
                rtl = True
        if level is not None and 0 <= level <= 8:
            outline.append({"level": level + 1, "text": text.encode()[:200].decode(errors="ignore"),
                            "paragraph": index})
        index += 1

    core = etree.fromstring(archive.read("docProps/core.xml"))
    title = core.findtext("{http://purl.org/dc/elements/1.1/}title") or ""
    creator = core.findtext("{http://purl.org/dc/elements/1.1/}creator") or ""
    return {
        "ok": True, "errorCode": "", "errorMessage": "", "kind": "docx",
        "sourceBytes": len(data), "features": [], "outline": outline,
        "properties": {"title": title, "creator": creator, "rightToLeft": rtl},
    }


def main() -> None:
    root = Path(__file__).resolve().parent
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "dev-fixtures"
    output.mkdir(parents=True, exist_ok=True)
    image_bytes = generated_png()
    builders = {
        "headings.docx": ("Quarterly field report", make_headings),
        "bullets.docx": ("Bullets and numbering", make_bullets),
        "tables.docx": ("Tables with headers", make_tables),
        "hyperlinks.docx": ("Hyperlinks", make_hyperlinks),
        "arabic.docx": ("تقرير", make_arabic),
        "images.docx": ("Images", lambda d: make_images(d, image_bytes)),
    }
    reports = {}
    for name, (title, build) in builders.items():
        document = new_document(title)
        build(document)
        buffer = io.BytesIO()
        document.save(buffer)
        data = deterministic_zip(buffer.getvalue())
        (output / name).write_bytes(data)
        reports[name] = static_report(data)
        print(f"wrote {name} ({len(data)} bytes)")
        if name == "images.docx":
            # The same document with its picture part removed: the image
            # relationship points at nothing.
            missing = without_parts(data, lambda part: part.startswith("word/media/"))
            (output / "missing-image.docx").write_bytes(missing)
            reports["missing-image.docx"] = static_report(missing)
            print(f"wrote missing-image.docx ({len(missing)} bytes)")
    (output / "reports.json").write_text(
        json.dumps(reports, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
