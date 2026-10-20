import sys, random
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
out = sys.argv[1]
random.seed(7)
WORDS = "the quick review of quarterly results shows steady progress across every regional team while costs remained within the approved budget and delivery dates held".split()
def para(n): return " ".join(random.choice(WORDS) for _ in range(n)).capitalize() + "."
def build(pages, name):
    d = Document()
    d.add_heading("Quarterly operations report", 0)
    for p in range(pages):
        d.add_heading(f"Section {p+1}: Regional summary", 1)
        d.add_paragraph(para(90))
        d.add_heading(f"{p+1}.1 Findings", 2)
        for i in range(3): d.add_paragraph(para(12), style="List Number" if i%2 else "List Bullet")
        r = d.add_paragraph().add_run(para(40)); r.bold = (p%3==0); r.font.color.rgb = RGBColor(0x1d,0x4e,0xd8)
        if p % 4 == 0:
            t = d.add_table(rows=5, cols=4); t.style = "Light Grid Accent 1"
            for ri,row in enumerate(t.rows):
                for ci,c in enumerate(row.cells): c.text = f"R{ri}C{ci} {random.randint(10,999)}"
            t.cell(0,0).merge(t.cell(0,1))
        d.add_paragraph(para(70))
        if p < pages-1: d.add_page_break()
    sec = d.sections[0]; sec.header.paragraphs[0].text = "LeanDoc spike header"; sec.footer.paragraphs[0].text = "Confidential footer"
    d.save(f"{out}/{name}")
build(3, "typical-3p.docx")
build(300, "long-300p.docx")
# Arabic RTL
d = Document()
h = d.add_heading("تقرير الربع الثالث", 1)
for _ in range(6):
    p = d.add_paragraph("هذا نص تجريبي باللغة العربية لاختبار الاتجاه من اليمين إلى اليسار، مع رقم 2026 وكلمة English في الوسط.")
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:bidi"); pPr.append(b)
    for r in p.runs:
        rtl = OxmlElement("w:rtl"); r._r.get_or_add_rPr().append(rtl)
d.save(f"{out}/arabic-rtl.docx")
print("ok")
